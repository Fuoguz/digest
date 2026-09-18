from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT_DIR = Path(__file__).resolve().parent
DOCX_PATH = OUT_DIR / "沉淀Digest_大学生创新创业训练计划创业训练项目计划书.docx"


PROJECT_NAME = "“沉淀 Digest”——面向高校学生的 AI 知识内化与自主学习能力提升平台"
DOC_TITLE = "大学生创新创业训练计划创业训练项目计划书"


def set_run_font(run, east="宋体", ascii_font="Times New Roman", size=None, bold=None, color=None):
    run.font.name = ascii_font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def set_para_font(paragraph, east="宋体", ascii_font="Times New Roman", size=12):
    for run in paragraph.runs:
        set_run_font(run, east=east, ascii_font=ascii_font, size=size)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, fill=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(str(text))
    set_run_font(run, east="宋体", size=10.5, bold=bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    if fill:
        set_cell_shading(cell, fill)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_text(hdr[i], h, bold=True, fill="D9EAF7", align=WD_ALIGN_PARAGRAPH.CENTER)
        if widths:
            hdr[i].width = Cm(widths[i])
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
            if widths:
                cells[i].width = Cm(widths[i])
    doc.add_paragraph()
    return table


def add_heading(doc, text, level=1):
    if level == 1:
        p = doc.add_paragraph()
        p.style = "Heading 1"
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(text)
        set_run_font(run, east="黑体", ascii_font="Arial", size=16, bold=True, color=(0, 47, 167))
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(8)
    else:
        p = doc.add_paragraph()
        p.style = "Heading 2"
        run = p.add_run(text)
        set_run_font(run, east="黑体", ascii_font="Arial", size=13, bold=True, color=(0, 0, 0))
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
    return p


def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, east="宋体", size=12)
    return p


def add_point(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.left_indent = Pt(0)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    set_run_font(run, east="宋体", size=12)
    return p


def add_footer_page_number(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("第 ")
    set_run_font(run, east="宋体", size=9)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    run2 = p.add_run(" 页")
    set_run_font(run2, east="宋体", size=9)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.6)
    section.right_margin = Cm(2.4)
    section.header_distance = Cm(1.4)
    section.footer_distance = Cm(1.2)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run(DOC_TITLE)
    set_run_font(run, east="宋体", size=9, color=(100, 100, 100))
    add_footer_page_number(section)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)


def add_cover(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(32)
    r = p.add_run("上海政法学院\n大学生创新创业训练计划项目")
    set_run_font(r, east="黑体", ascii_font="Arial", size=18, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(36)
    r = p.add_run(DOC_TITLE)
    set_run_font(r, east="黑体", ascii_font="Arial", size=22, bold=True, color=(0, 47, 167))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(24)
    r = p.add_run(PROJECT_NAME)
    set_run_font(r, east="黑体", ascii_font="Arial", size=16, bold=True)

    table = doc.add_table(rows=6, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    info = [
        ("项目类型", "创业训练项目"),
        ("项目负责人", "张少毅"),
        ("所在学院", "上海政法学院上海纪录片学院"),
        ("专业", "网络与新媒体"),
        ("指导方向", "AI+教育、教育数字化、自主学习能力提升"),
        ("填报时间", "2026 年 5 月"),
    ]
    for i, (k, v) in enumerate(info):
        set_cell_text(table.cell(i, 0), k, bold=True, fill="EAF2F8", align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(table.cell(i, 1), v)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(42)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("项目计划书")
    set_run_font(r, east="宋体", size=12)
    doc.add_page_break()


def add_contents(doc):
    add_heading(doc, "目录", 1)
    items = [
        "一、项目概述",
        "二、项目背景与问题提出",
        "三、目标用户与需求分析",
        "四、产品定位与核心功能",
        "五、产品原型与技术实现路径",
        "六、市场分析",
        "七、竞品分析",
        "八、商业模式设计",
        "九、市场验证与用户调研计划",
        "十、模拟运营方案",
        "十一、团队成员与分工",
        "十二、项目实施计划",
        "十三、经费预算与使用说明",
        "十四、风险分析与应对措施",
        "十五、预期成果",
        "十六、后续发展规划",
    ]
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.5
        run = p.add_run(item)
        set_run_font(run, east="宋体", size=12)
    doc.add_page_break()


def build_doc():
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_contents(doc)

    add_heading(doc, "一、项目概述", 1)
    add_heading(doc, "（一）项目简介", 2)
    add_point(doc, "1. “沉淀 Digest”是一款面向高校学生的 AI 知识内化与自主学习能力提升平台，主要服务于课程学习、论文阅读、资料整理、考试复习、课堂汇报和项目研究等学习场景。")
    add_point(doc, "2. 项目通过 AI 结构化摘要、复习问题生成、知识节点沉淀和知识图谱展示等功能，帮助学生将一次性阅读内容转化为可积累、可复习、可复用的个人知识资产。")
    add_point(doc, "3. 本项目属于创业训练项目，重点围绕产品原型开发、用户调研、市场分析、竞品分析、商业模式设计、模拟运营和创业计划完善等环节开展。")
    add_heading(doc, "（二）项目定位", 2)
    add_point(doc, "1. 本项目不定位为普通 AI 总结工具，而定位为面向高校学生的 AI 知识内化平台。")
    add_point(doc, "2. 项目重点解决学生“看过资料但难以真正理解、记忆和复用”的问题。")
    add_point(doc, "3. 项目核心价值在于将 AI 摘要能力与主动回忆、知识管理、长期复习和学习场景结合，形成完整学习闭环。")
    add_heading(doc, "（三）已有基础", 2)
    add_point(doc, "1. 项目已有 Digest 前端网页原型和 GitHub 项目基础。")
    add_point(doc, "2. 项目负责人具备产品策划、网页原型开发、AI 辅助开发和项目表达能力。")
    add_point(doc, "3. 项目负责人具备 Python、JavaScript 等基础，能够使用 Codex、Claude Code 等 AI 辅助开发工具进行代码生成、页面修改、功能调试和产品迭代。")
    add_point(doc, "4. 团队可依托上海政法学院、上海纪录片学院、网络与新媒体专业背景、校内学生群体、班级社群和课程学习场景开展问卷调查、用户访谈和小范围试用。")

    add_heading(doc, "二、项目背景与问题提出", 1)
    add_heading(doc, "（一）高校学生学习资料处理压力增加", 2)
    add_point(doc, "1. 随着教育数字化发展，高校学生接触的学习资料不断增加，包括课程课件、教材章节、论文文献、网页内容、政策文本、案例材料和项目调研资料等。")
    add_point(doc, "2. 学生获取资料更加便利，但也面临资料过多、重点难提炼、阅读后容易遗忘、笔记分散、资料难以复用等问题。")
    add_point(doc, "3. 在课程学习、论文写作、考试复习和课堂汇报中，学生需要的不只是快速获取信息，更需要将信息转化为可理解、可记忆、可复用的知识结构。")
    add_heading(doc, "（二）现有工具存在不足", 2)
    add_point(doc, "1. ChatGPT、Kimi、豆包等通用 AI 工具可以完成文本总结和问答，但多数停留在一次性回答层面。")
    add_point(doc, "2. Notion AI、Obsidian 等工具具备笔记和知识管理能力，但对普通学生存在一定使用门槛。")
    add_point(doc, "3. 飞书妙记、秘塔 AI 搜索、夸克 AI 等工具更偏向会议记录、搜索和信息整理，不完全适配高校学生长期学习闭环。")
    add_heading(doc, "（三）项目提出的必要性", 2)
    add_point(doc, "1. 高校学生需要一款低门槛、轻量化、可持续使用的 AI 学习辅助产品。")
    add_point(doc, "2. Digest 通过“内容输入—AI 解析—复习问题—知识节点—知识图谱—后续复习”的流程，帮助学生形成知识内化闭环。")
    add_point(doc, "3. 项目具有 AI+教育、教育数字化、新文科实践和自主学习能力提升等方面的应用价值。")

    add_heading(doc, "三、目标用户与需求分析", 1)
    add_heading(doc, "（一）目标用户", 2)
    add_point(doc, "1. 本项目主要面向高校本科生和研究生。")
    add_point(doc, "2. 初期重点面向上海政法学院学生，尤其是网络与新媒体、新闻传播、法学、社会学、公共管理等专业学生。")
    add_point(doc, "3. 后续可逐步扩展至更多高校学生群体。")
    add_heading(doc, "（二）核心学习场景", 2)
    add_table(
        doc,
        ["场景", "用户需求", "Digest 对应功能"],
        [
            ["课程学习", "整理课件、教材章节、课堂资料", "结构化摘要、重点提炼、知识节点"],
            ["论文阅读", "提炼研究问题、核心观点和论据", "AI 摘要、论据整理、复习问题"],
            ["考试复习", "形成复习提纲和自测问题", "复习问题生成、知识节点沉淀"],
            ["课堂汇报", "快速整理资料结构和表达框架", "摘要导出、观点整理、素材复用"],
            ["项目研究", "管理调研资料、案例材料和竞品资料", "项目知识节点、图谱展示、团队资料管理"],
        ],
        [3.0, 6.0, 6.0],
    )
    add_heading(doc, "（三）用户需求", 2)
    add_point(doc, "1. 快速理解长文本内容，提炼核心观点和结构。")
    add_point(doc, "2. 将阅读内容转化为复习问题和思考题。")
    add_point(doc, "3. 保存学习资料，形成长期可积累的知识节点。")
    add_point(doc, "4. 可视化查看不同知识之间的联系。")
    add_point(doc, "5. 在论文写作、课堂汇报和项目研究中复用已整理资料。")

    add_heading(doc, "四、产品定位与核心功能", 1)
    add_heading(doc, "（一）产品定位", 2)
    add_point(doc, "1. Digest 是面向高校学生的 AI 知识内化平台。")
    add_point(doc, "2. Digest 是服务课程学习、论文阅读、资料整理和自主学习能力提升的 AI 学习辅助平台。")
    add_point(doc, "3. Digest 是帮助学生建立个人知识资产的轻量化知识管理工具。")
    add_heading(doc, "（二）核心功能", 2)
    add_table(
        doc,
        ["功能模块", "功能说明", "应用价值"],
        [
            ["内容输入", "支持粘贴文章、网页内容、课程资料、Markdown 文本", "统一学习资料入口"],
            ["AI 结构化摘要", "自动提炼核心观点、文章结构和论据", "提升阅读理解效率"],
            ["复习问题生成", "基于主动回忆机制生成复习题和思考题", "提升知识留存率"],
            ["知识节点沉淀", "将一次阅读内容保存为知识节点", "形成个人知识库"],
            ["知识图谱展示", "可视化展示知识之间的关联", "帮助理解和复用知识"],
            ["导出与分享", "后续支持导出笔记、复习卡片和汇报素材", "服务复习、写作和展示"],
        ],
        [3.2, 6.4, 5.4],
    )
    add_heading(doc, "（三）后续拓展方向", 2)
    for i, text in enumerate(["学习路径推荐。", "间隔复习提醒。", "法学案例材料整理。", "政策文本解读。", "新文科课程学习辅助。", "面向学生团队的项目资料管理。"], 1):
        add_point(doc, f"{i}. {text}")

    add_heading(doc, "五、产品原型与技术实现路径", 1)
    add_heading(doc, "（一）产品原型基础", 2)
    add_point(doc, "1. 项目已有 Digest 前端网页原型。")
    add_point(doc, "2. 项目已有 GitHub 项目基础，便于代码托管、版本管理和后续展示。")
    add_point(doc, "3. 当前原型已具备继续开发、演示和迭代的基础。")
    add_heading(doc, "（二）技术实现路径", 2)
    add_point(doc, "1. 前端网页开发：使用 HTML、CSS、JavaScript 构建内容输入、结果展示、知识节点和知识图谱页面。")
    add_point(doc, "2. AI API 调用：调用现有成熟 AI 模型 API，实现结构化摘要、观点提炼、论据整理和复习问题生成。")
    add_point(doc, "3. 本地存储：初期采用浏览器本地存储保存知识节点，降低服务器成本。")
    add_point(doc, "4. 知识图谱可视化：使用 D3.js 等前端可视化工具展示知识节点关系。")
    add_point(doc, "5. 低成本部署：通过 GitHub Pages、Vercel 等平台部署网页原型，便于校内试用和答辩展示。")
    add_heading(doc, "（三）技术边界", 2)
    add_point(doc, "1. 项目不进行自研大模型训练。")
    add_point(doc, "2. 项目不以复杂算法研发为目标。")
    add_point(doc, "3. 项目重点在于将现有 AI 能力与高校学习场景结合，完成可用产品原型和创业验证。")

    add_heading(doc, "六、市场分析", 1)
    add_heading(doc, "（一）市场需求", 2)
    add_point(doc, "1. 高校学生在学习过程中存在大量资料整理和文本处理需求。")
    add_point(doc, "2. AI 工具已逐渐进入学生学习场景，但多数工具缺少长期知识沉淀机制。")
    add_point(doc, "3. 学生对低门槛、高效率、可复用的学习辅助工具具有潜在需求。")
    add_heading(doc, "（二）校园市场特点", 2)
    add_point(doc, "1. 用户集中，便于调研和试用。")
    add_point(doc, "2. 学习场景明确，便于设计具体功能。")
    add_point(doc, "3. 反馈周期较短，适合创业训练项目进行迭代验证。")
    add_heading(doc, "（三）商业机会", 2)
    add_point(doc, "1. 基础功能免费可降低学生使用门槛。")
    add_point(doc, "2. 高阶功能订阅可探索持续变现。")
    add_point(doc, "3. 课程复习包、资料整理服务和团队项目服务具有场景化付费可能。")
    add_point(doc, "4. 校园学习社群可作为用户留存和后续运营入口。")

    add_heading(doc, "七、竞品分析", 1)
    add_table(
        doc,
        ["竞品", "主要优势", "主要不足", "Digest 差异化方向"],
        [
            ["ChatGPT", "通用问答和文本生成能力强", "缺少默认学习闭环", "聚焦高校学习场景和知识沉淀"],
            ["Kimi", "长文本处理能力较强", "偏一次性总结和问答", "增加复习问题和节点积累"],
            ["豆包", "使用门槛低，适合日常问答", "学习资料管理能力有限", "强化课程、论文、复习场景"],
            ["Notion AI", "文档管理和 AI 写作结合", "学习成本较高", "提供更轻量的学生学习工具"],
            ["Obsidian", "知识管理和双链能力强", "依赖手动维护，门槛较高", "用 AI 降低知识管理门槛"],
            ["飞书妙记", "会议记录和协作能力强", "不专注个人学习内化", "聚焦阅读、复习和知识节点"],
            ["秘塔 AI 搜索", "搜索和资料整合能力较强", "偏信息获取", "强化长期复习和知识复用"],
            ["夸克 AI", "搜索和文档处理便利", "学习闭环不足", "建立知识节点和图谱展示"],
        ],
        [2.2, 4.0, 4.0, 4.8],
    )

    add_heading(doc, "八、商业模式设计", 1)
    add_heading(doc, "（一）免费入口", 2)
    add_point(doc, "1. 提供基础文本输入、AI 摘要、少量复习问题生成和有限知识节点保存功能。")
    add_point(doc, "2. 降低学生首次使用门槛。")
    add_point(doc, "3. 通过免费功能积累真实用户反馈，验证产品价值。")
    add_heading(doc, "（二）高阶订阅", 2)
    add_table(
        doc,
        ["付费功能", "说明"],
        [
            ["更多知识节点额度", "支持长期保存更多学习内容"],
            ["更长文本处理", "支持论文、政策文本、长篇资料处理"],
            ["批量导出", "支持导出 Markdown、Word、复习卡片"],
            ["知识图谱增强", "支持更多节点关系和筛选能力"],
            ["复习计划", "支持间隔复习提醒和学习计划管理"],
            ["多端同步", "后续支持跨设备学习资料管理"],
        ],
        [5.0, 10.0],
    )
    add_heading(doc, "（三）资料整理服务", 2)
    add_point(doc, "1. 面向学生课程学习和论文阅读，提供学习资料结构化整理服务。")
    add_point(doc, "2. 可将学生自有资料整理为摘要、复习问题、知识节点和汇报提纲。")
    add_point(doc, "3. 服务形式可按单次资料包、课程模块或项目材料进行设计。")
    add_heading(doc, "（四）课程复习包", 2)
    add_point(doc, "1. 面向期末复习、课程汇报和课程论文需求，探索课程复习包服务。")
    add_point(doc, "2. 复习包内容包括课程重点、概念卡片、复习问题和知识节点。")
    add_point(doc, "3. 该服务应在合规前提下基于学生自有资料或公开资料开展。")
    add_heading(doc, "（五）团队项目服务", 2)
    add_point(doc, "1. 面向大创项目、课程小组作业、社会调查和创新创业赛事团队。")
    add_point(doc, "2. 提供调研资料、访谈记录、竞品资料和答辩素材的结构化管理。")
    add_point(doc, "3. 后续可探索按项目周期或团队人数进行轻量收费。")
    add_heading(doc, "（六）校园学习社群", 2)
    add_point(doc, "1. 建立围绕 AI 学习方法、论文阅读、资料整理和复习策略的学习社群。")
    add_point(doc, "2. 提供模板分享、工具使用指导、学习案例和线上交流。")
    add_point(doc, "3. 社群可作为产品反馈、用户留存和后续服务转化渠道。")

    add_heading(doc, "九、市场验证与用户调研计划", 1)
    add_heading(doc, "（一）问卷调研", 2)
    add_point(doc, "1. 面向高校学生设计问卷，了解学习资料整理、论文阅读、考试复习和 AI 工具使用情况。")
    add_point(doc, "2. 调研学生对结构化摘要、复习问题、知识节点和知识图谱功能的需求。")
    add_point(doc, "3. 调研学生对免费功能、高阶订阅和资料整理服务的接受程度。")
    add_heading(doc, "（二）用户访谈", 2)
    add_point(doc, "1. 访谈不同专业、不同年级学生的学习资料处理习惯。")
    add_point(doc, "2. 了解学生在课程学习、论文阅读、课堂汇报和项目研究中的真实痛点。")
    add_point(doc, "3. 收集用户对 Digest 原型界面、功能流程和输出结果的反馈。")
    add_heading(doc, "（三）校内试用", 2)
    add_point(doc, "1. 邀请校内学生使用 Digest 处理课程资料、论文片段或项目材料。")
    add_point(doc, "2. 观察用户是否能理解产品流程，是否愿意持续使用。")
    add_point(doc, "3. 根据反馈优化摘要结构、复习问题质量、知识节点展示和图谱交互。")
    add_heading(doc, "（四）付费意愿验证", 2)
    add_point(doc, "1. 通过问卷和访谈了解学生对高阶订阅的接受程度。")
    add_point(doc, "2. 验证课程复习包、资料整理服务、团队项目服务和学习社群的可行性。")
    add_point(doc, "3. 形成商业模式验证结论，为后续创业计划书完善提供依据。")

    add_heading(doc, "十、模拟运营方案", 1)
    add_heading(doc, "（一）产品展示", 2)
    add_point(doc, "1. 建立 GitHub 项目页面和网页演示入口。")
    add_point(doc, "2. 制作产品演示视频和项目介绍材料。")
    add_point(doc, "3. 在答辩、课程展示和校内交流中展示产品原型。")
    add_heading(doc, "（二）种子用户获取", 2)
    add_point(doc, "1. 依托班级社群、课程小组、学生组织和学院渠道邀请学生试用。")
    add_point(doc, "2. 初期重点获取真实反馈，不追求大规模用户数量。")
    add_point(doc, "3. 记录用户使用场景、问题反馈和改进建议。")
    add_heading(doc, "（三）内容运营", 2)
    add_point(doc, "1. 设计“用 Digest 整理一篇论文”的示例内容。")
    add_point(doc, "2. 设计“用 Digest 生成课程复习问题”的使用案例。")
    add_point(doc, "3. 设计“用 Digest 沉淀项目调研资料”的项目场景案例。")
    add_point(doc, "4. 通过微信公众号、小红书、B站、GitHub 页面或校内展示活动发布使用教程和项目进展。")
    add_heading(doc, "（四）模拟付费方案测试", 2)
    add_table(
        doc,
        ["方案", "验证内容"],
        [
            ["免费版与高阶版", "学生是否愿意为更多节点、导出和复习计划付费"],
            ["资料整理服务", "学生是否愿意为课程资料结构化整理付费"],
            ["课程复习包", "学生是否接受按课程或主题生成复习包"],
            ["团队项目服务", "学生团队是否需要项目资料管理工具"],
            ["校园学习社群", "学生是否愿意加入 AI 学习方法社群"],
        ],
        [5.0, 10.0],
    )

    add_heading(doc, "十一、团队成员与分工", 1)
    add_table(
        doc,
        ["成员", "角色", "主要分工"],
        [
            ["张少毅", "项目负责人", "项目统筹、产品定位、网页原型开发、核心材料撰写、商业模式设计、答辩汇报、进度管理"],
            ["成员 A", "用户调研负责人", "问卷设计、问卷分发、用户访谈、数据整理、用户需求分析"],
            ["成员 B", "市场与展示负责人", "竞品分析、市场资料收集、视觉展示、PPT、宣传材料、演示视频辅助"],
        ],
        [3.0, 3.8, 8.2],
    )

    add_heading(doc, "十二、项目实施计划", 1)
    add_table(
        doc,
        ["阶段", "时间", "主要任务", "阶段成果"],
        [
            ["资料查阅与项目定位", "2026 年 6 月—2026 年 7 月", "查阅 AI+教育、知识管理、主动回忆、教育数字化等资料", "项目定位与资料整理"],
            ["产品方案设计", "2026 年 7 月—2026 年 8 月", "明确目标用户、功能模块、技术路线和商业模式假设", "初步产品方案"],
            ["原型优化与开题准备", "2026 年 8 月—2026 年 9 月", "优化网页原型，准备开题材料", "开题材料、原型演示"],
            ["用户调研与竞品分析", "2026 年 9 月—2026 年 11 月", "开展问卷、访谈、竞品体验", "用户调研材料、竞品分析"],
            ["校内试用与中期检查", "2026 年 11 月—2026 年 12 月", "开展小范围试用，收集反馈", "中期汇报材料"],
            ["反馈整理与产品迭代", "2026 年 12 月—2027 年 1 月", "分析反馈，优化核心功能", "产品迭代版本"],
            ["创业材料撰写", "2027 年 1 月—2027 年 3 月", "完成创业计划书、调研报告、模拟运营报告", "项目成果材料"],
            ["结项答辩与成果推广", "2027 年 3 月—2027 年 4 月", "制作答辩 PPT、演示视频、结题材料", "结项答辩材料"],
        ],
        [3.2, 3.5, 5.5, 3.8],
    )

    add_heading(doc, "十三、经费预算与使用说明", 1)
    add_table(
        doc,
        ["经费用途", "金额", "使用说明"],
        [
            ["产品开发与技术服务", "4000 元", "用于网页原型优化、功能模块开发、前端交互完善等"],
            ["AI API 调用与服务器部署", "3000 元", "用于 AI 模型接口调用、网页部署、域名或服务器等"],
            ["用户调研与访谈", "2000 元", "用于问卷激励、访谈记录、样本收集、调研材料整理等"],
            ["UI 设计与视觉素材", "2000 元", "用于界面优化、图标素材、展示图、宣传物料设计等"],
            ["产品测试与校园推广", "2000 元", "用于校内试用、用户反馈收集、校园推广和活动组织等"],
            ["答辩展示与视频制作", "1500 元", "用于项目 PPT、演示视频、宣传短片和结项展示材料制作"],
            ["资料购置与其他支出", "500 元", "用于相关书籍、资料、打印和其他必要支出"],
            ["合计", "15000 元", "上海政法学院大学生创新创业训练计划项目资助经费"],
        ],
        [4.2, 2.6, 8.0],
    )

    add_heading(doc, "十四、风险分析与应对措施", 1)
    add_table(
        doc,
        ["风险类型", "具体表现", "应对措施"],
        [
            ["需求验证不足", "学生可能认为通用 AI 工具已能满足需求", "通过具体课程、论文、复习和项目场景验证真实需求"],
            ["用户持续使用不足", "用户初次体验后可能缺少持续使用动力", "强化知识节点、复习问题和图谱沉淀功能"],
            ["商业变现不确定", "学生付费能力有限", "采用免费入口和轻量付费，重点验证付费意愿"],
            ["技术实现风险", "AI API 成本、输出质量和图谱展示可能存在不稳定", "控制功能边界，采用成熟 API 和轻量前端方案"],
            ["数据隐私风险", "用户输入资料可能涉及个人作业或课程材料", "初期采用本地存储，设置隐私提示"],
            ["团队执行风险", "成员受课程、实习、考试影响", "制定阶段计划，明确分工，优先完成核心成果"],
        ],
        [3.5, 5.0, 6.2],
    )

    add_heading(doc, "十五、预期成果", 1)
    add_table(
        doc,
        ["成果类型", "具体内容"],
        [
            ["产品原型", "完成 Digest AI 知识内化平台网页原型一套"],
            ["用户调研报告", "形成问卷设计、访谈记录、用户痛点分析和功能需求分析"],
            ["竞品分析报告", "对 ChatGPT、Kimi、豆包、Notion AI、Obsidian、飞书妙记等产品进行分析"],
            ["创业计划书", "完成项目背景、市场分析、商业模式、运营方案和风险分析"],
            ["模拟运营报告", "记录校内试用、用户反馈、内容传播和付费意愿验证情况"],
            ["答辩材料", "制作项目答辩 PPT、产品演示视频和宣传材料"],
            ["后续参赛基础", "为中国国际大学生创新大赛、“挑战杯”、知行杯等赛事准备项目材料"],
        ],
        [4.0, 11.0],
    )

    add_heading(doc, "十六、后续发展规划", 1)
    add_heading(doc, "（一）产品功能优化", 2)
    add_point(doc, "1. 优化 AI 结构化摘要质量。")
    add_point(doc, "2. 提升复习问题生成效果。")
    add_point(doc, "3. 完善知识节点管理和知识图谱交互。")
    add_point(doc, "4. 增加导出、复习提醒和学习路径推荐功能。")
    add_heading(doc, "（二）学习场景拓展", 2)
    add_point(doc, "1. 拓展课程复习场景。")
    add_point(doc, "2. 拓展论文阅读场景。")
    add_point(doc, "3. 拓展课堂汇报场景。")
    add_point(doc, "4. 拓展法学案例材料整理和政策文本解读场景。")
    add_point(doc, "5. 拓展新文科课程学习辅助场景。")
    add_heading(doc, "（三）商业模式深化", 2)
    add_point(doc, "1. 持续验证高阶订阅模式。")
    add_point(doc, "2. 探索资料整理服务和课程复习包服务。")
    add_point(doc, "3. 探索学生团队项目资料管理服务。")
    add_point(doc, "4. 建立校园学习社群，提高用户留存和反馈效率。")
    add_heading(doc, "（四）成果推广与赛事申报", 2)
    add_point(doc, "1. 通过校内展示、课程汇报和学院活动推广项目成果。")
    add_point(doc, "2. 通过 GitHub 项目页面和新媒体平台展示产品进展。")
    add_point(doc, "3. 后续申报中国国际大学生创新大赛、“挑战杯”、知行杯等赛事。")
    add_point(doc, "4. 根据赛事和项目发展需要，继续完善商业计划书、路演 PPT 和演示视频。")

    doc.core_properties.title = DOC_TITLE
    doc.core_properties.subject = PROJECT_NAME
    doc.core_properties.author = "张少毅"
    doc.save(DOCX_PATH)
    return DOCX_PATH


if __name__ == "__main__":
    print(build_doc())
